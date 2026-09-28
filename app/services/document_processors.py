"""Characteristic-driven, replaceable document-processing adapters."""

from __future__ import annotations

import io
from dataclasses import dataclass
from typing import Protocol

import fitz
from docx import Document as WordDocument
from openpyxl import load_workbook
from PIL import Image

from app.models.understanding import BlockType, ContentBlock, PageRepresentation


class ProcessorUnavailable(RuntimeError):
    pass


class OCRProcessor(Protocol):
    name: str
    version: str

    def extract(self, image: Image.Image, *, page_number: int) -> list[ContentBlock]: ...


@dataclass(frozen=True, slots=True)
class ProcessorResult:
    pages: tuple[PageRepresentation, ...]
    blocks: tuple[ContentBlock, ...]


class PDFProcessor:
    name = "pymupdf"
    version = fitz.VersionBind

    def process(self, data: bytes) -> ProcessorResult:
        try:
            document = fitz.open(stream=data, filetype="pdf")
            pages: list[PageRepresentation] = []
            blocks: list[ContentBlock] = []
            for number, page in enumerate(document, 1):
                page_blocks = []
                for order, item in enumerate(page.get_text("blocks")):
                    x0, y0, x1, y1, text = item[:5]
                    if text.strip():
                        page_blocks.append(
                            ContentBlock(
                                type=BlockType.TEXT,
                                text=text.strip(),
                                page_number=number,
                                reading_order=order,
                                bbox=(x0, y0, x1, y1),
                            )
                        )
                pages.append(
                    PageRepresentation(
                        number, page.rect.width, page.rect.height, tuple(page_blocks)
                    )
                )
                blocks.extend(page_blocks)
            document.close()
            if not blocks:
                raise ProcessorUnavailable("OCR processor is required for this scanned PDF")
            return ProcessorResult(tuple(pages), tuple(blocks))
        except ProcessorUnavailable:
            raise
        except Exception as exc:
            raise ValueError("invalid PDF document") from exc


class ImageOCRProcessor:
    name = "ocr"

    def __init__(self, ocr: OCRProcessor | None = None) -> None:
        self.ocr = ocr

    def process(self, data: bytes) -> ProcessorResult:
        try:
            image = Image.open(io.BytesIO(data))
            image.load()
        except Exception as exc:
            raise ValueError("invalid image document") from exc
        if self.ocr is None:
            raise ProcessorUnavailable("OCR processor is not configured")
        blocks = tuple(self.ocr.extract(image, page_number=1))
        return ProcessorResult((PageRepresentation(1, image.width, image.height, blocks),), blocks)


class DOCXProcessor:
    name = "python-docx"
    version = "1"

    def process(self, data: bytes) -> ProcessorResult:
        try:
            document = WordDocument(io.BytesIO(data))
            blocks = tuple(
                ContentBlock(type=BlockType.TEXT, text=p.text.strip(), reading_order=i)
                for i, p in enumerate(document.paragraphs)
                if p.text.strip()
            )
            return ProcessorResult((PageRepresentation(1, blocks=blocks),), blocks)
        except Exception as exc:
            raise ValueError("invalid DOCX document") from exc


class XLSXProcessor:
    name = "openpyxl"
    version = "1"

    def process(self, data: bytes) -> ProcessorResult:
        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            pages: list[PageRepresentation] = []
            blocks: list[ContentBlock] = []
            for page_number, sheet in enumerate(workbook.worksheets, 1):
                sheet_blocks = []
                for order, row in enumerate(sheet.iter_rows(values_only=True)):
                    text = " | ".join(str(value) for value in row if value is not None).strip()
                    if text:
                        sheet_blocks.append(
                            ContentBlock(
                                type=BlockType.TEXT,
                                text=text,
                                page_number=page_number,
                                reading_order=order,
                            )
                        )
                pages.append(PageRepresentation(page_number, blocks=tuple(sheet_blocks)))
                blocks.extend(sheet_blocks)
            workbook.close()
            return ProcessorResult(tuple(pages), tuple(blocks))
        except Exception as exc:
            raise ValueError("invalid XLSX document") from exc


__all__ = [
    "DOCXProcessor",
    "ImageOCRProcessor",
    "OCRProcessor",
    "PDFProcessor",
    "ProcessorResult",
    "ProcessorUnavailable",
    "XLSXProcessor",
]
